"""Bulk import of customers and suppliers from a spreadsheet.

A company moving to Vezano has its ledger in Excel: hundreds of names, phones
and balances. Two steps, both through the same endpoint: a dry run that
returns every row with what would happen to it (created, updated, skipped,
or the error), then the real run. Rows match an existing party by phone
first, then by exact name, so re-importing the same sheet updates rather
than duplicates. An opening balance column becomes the same document an
opening balance from the screen does (see core.opening_balances), and only
for a party that has none — an approver-only step, like the screen.

Accepted: .xlsx (openpyxl) or .csv (UTF-8, comma or semicolon). Header names
are matched loosely in Arabic or English: الاسم/name, الهاتف/phone,
البريد/email, العنوان/address, الرصيد الافتتاحي/opening balance,
مهلة السداد/payment terms.
"""

import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.translation import gettext as _
from rest_framework.exceptions import ValidationError

MAX_ROWS = 2000

# canonical field -> header spellings (lowercased, stripped)
HEADERS = {
    "name": {"name", "الاسم", "اسم", "customer", "supplier", "العميل", "المورد", "الاسم الكامل"},
    "phone": {"phone", "mobile", "tel", "الهاتف", "هاتف", "الجوال", "جوال", "رقم الهاتف"},
    "email": {"email", "e-mail", "البريد", "البريد الإلكتروني", "الايميل", "الإيميل"},
    "address": {"address", "العنوان", "عنوان"},
    "opening_balance": {
        "opening balance", "opening_balance", "balance", "الرصيد الافتتاحي", "رصيد افتتاحي",
        "الرصيد", "رصيد", "المستحق", "الدين",
    },
    "payment_terms_days": {
        "payment terms", "terms", "payment_terms_days", "مهلة السداد", "مهلة السداد (بالأيام)",
        "شروط الدفع", "أيام السداد",
    },
}
TEMPLATE_COLUMNS = ["name", "phone", "email", "address", "opening_balance", "payment_terms_days"]


def _norm_header(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower().replace("ـ", "")


def _map_headers(headers):
    mapping = {}
    for index, raw in enumerate(headers):
        key = _norm_header(raw)
        for field, spellings in HEADERS.items():
            if key in spellings and field not in mapping:
                mapping[field] = index
                break
    if "name" not in mapping:
        raise ValidationError(
            {"file": [_("The first row must name the columns, and one of them must be the name.")]}
        )
    return mapping


def _rows_from_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - dependency is in requirements
        raise ValidationError({"file": [_("Excel import is not available on this server.")]})
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    for row in sheet.iter_rows(values_only=True):
        yield ["" if v is None else v for v in row]


def _rows_from_csv(data):
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        yield row


def parse_sheet(uploaded):
    """Yield (row_number, {field: value}) for every data row."""
    name = (getattr(uploaded, "name", "") or "").lower()
    data = uploaded.read()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        rows = _rows_from_xlsx(data)
    elif name.endswith(".csv") or name.endswith(".txt") or not name:
        rows = _rows_from_csv(data)
    else:
        raise ValidationError({"file": [_("Upload an .xlsx or .csv file.")]})
    rows = iter(rows)
    try:
        headers = next(rows)
    except StopIteration:
        raise ValidationError({"file": [_("The file is empty.")]})
    mapping = _map_headers(headers)
    count = 0
    for number, row in enumerate(rows, start=2):
        if not any(str(v).strip() for v in row):
            continue
        count += 1
        if count > MAX_ROWS:
            raise ValidationError(
                {"file": [_("At most %(n)s rows per file.") % {"n": MAX_ROWS}]}
            )
        record = {}
        for field, index in mapping.items():
            value = row[index] if index < len(row) else ""
            record[field] = str(value).strip() if value is not None else ""
        yield number, record


def _clean_amount(value, field):
    if value in ("", None):
        return None
    text = str(value).replace(",", "").replace("٬", "").strip()
    text = text.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise ValueError(_("%(field)s is not a number.") % {"field": field})
    return amount


def _clean_row(record):
    name = record.get("name", "").strip()
    if not name:
        raise ValueError(_("Name is missing."))
    email = record.get("email", "").strip()
    if email and ("@" not in email or " " in email):
        raise ValueError(_("Email is not valid."))
    phone = re.sub(r"[^\d+]", "", record.get("phone", "")
                   .translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")))
    opening = _clean_amount(record.get("opening_balance"), _("Opening balance"))
    if opening is not None and opening < 0:
        raise ValueError(_("Opening balance cannot be negative."))
    terms = _clean_amount(record.get("payment_terms_days"), _("Payment terms"))
    if terms is not None and (terms < 0 or terms > 365 or terms != terms.to_integral_value()):
        raise ValueError(_("Payment terms must be whole days between 0 and 365."))
    return {
        "name": name[:255], "phone": phone[:64], "email": email[:254],
        "address": record.get("address", "").strip(),
        "opening_balance": opening,
        "payment_terms_days": int(terms) if terms is not None else None,
    }


def run_import(model, company, user, uploaded, *, dry_run, opening_balance_fn,
               supports_terms):
    """Returns {"summary": {...}, "rows": [...]}. Wrapped in a transaction that
    is rolled back on a dry run, so the preview reflects real validation."""
    from core.rbac import can_approve_high_value

    rows_out = []
    summary = {"created": 0, "updated": 0, "balances": 0, "errors": 0, "skipped": 0}
    seen_in_file = set()
    approver = can_approve_high_value(user)

    with transaction.atomic():
        for number, record in parse_sheet(uploaded):
            out = {"row": number, "name": record.get("name", ""), "phone": record.get("phone", "")}
            try:
                clean = _clean_row(record)
            except ValueError as exc:
                out.update(status="error", message=str(exc))
                summary["errors"] += 1
                rows_out.append(out)
                continue
            key = clean["phone"] or clean["name"].casefold()
            if key in seen_in_file:
                out.update(status="skipped", message=_("Duplicate of an earlier row in the file."))
                summary["skipped"] += 1
                rows_out.append(out)
                continue
            seen_in_file.add(key)

            existing = None
            if clean["phone"]:
                existing = model.objects.filter(company=company, phone=clean["phone"]).first()
            if existing is None:
                existing = model.objects.filter(company=company, name__iexact=clean["name"]).first()

            fields = {"name": clean["name"], "phone": clean["phone"], "email": clean["email"],
                      "address": clean["address"]}
            if supports_terms and clean["payment_terms_days"] is not None:
                fields["payment_terms_days"] = clean["payment_terms_days"]
            if existing is None:
                party = model.objects.create(company=company, **fields)
                out["status"] = "created"
                summary["created"] += 1
            else:
                for attr, value in fields.items():
                    if value:
                        setattr(existing, attr, value)
                existing.save()
                party = existing
                out["status"] = "updated"
                summary["updated"] += 1

            if clean["opening_balance"]:
                if not approver:
                    out["message"] = _(
                        "Opening balance ignored: only a manager or owner may record one."
                    )
                elif opening_balance_fn(party):
                    out["message"] = _("Opening balance kept: this account already has one.")
                else:
                    opening_balance_fn(party, user, clean["opening_balance"])
                    out["balance"] = str(clean["opening_balance"])
                    summary["balances"] += 1
            rows_out.append(out)

        if dry_run:
            transaction.set_rollback(True)
    return {"summary": summary, "dry_run": dry_run, "rows": rows_out}
