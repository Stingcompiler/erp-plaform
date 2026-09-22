"""A backup the owner can actually read: the same snapshot as a workbook or
a zip of CSVs, with names instead of ids and headers in their language."""
import io
import zipfile
from decimal import Decimal

from django.urls import reverse
from rest_framework.test import APITestCase

from accounts.models import Role, User
from inventory.models import Category, Product, Warehouse
from org.models import Company
from sales.models import Customer, Invoice


class BackupFormatTests(APITestCase):
    def setUp(self):
        self.company = Company.objects.create(name="Alpha Trading", currency="SDG")
        role = Role.objects.create(name="Business Owner", scope_level=Role.SCOPE_BUSINESS)
        self.user = User.objects.create_user(
            email="owner@alpha.test", password="passw0rd123",
            company=self.company, role=role,
        )
        warehouse = Warehouse.objects.create(company=self.company, name="المخزن الرئيسي")
        category = Category.objects.create(company=self.company, name="مواد غذائية")
        Product.objects.create(
            company=self.company, sku="SKU1", name="أرز بسمتي", category=category,
            cost_price=Decimal("6.50"), sale_price=Decimal("10"), barcode="0629000000001",
        )
        customer = Customer.objects.create(company=self.company, name="بقالة الحي")
        Invoice.objects.create(
            company=self.company, customer=customer, warehouse=warehouse, number=1,
            subtotal=Decimal("100"), total=Decimal("100"),
        )
        self.client.post(
            reverse("auth-login"),
            {"email": "owner@alpha.test", "password": "passw0rd123", "device_id": "TEST"},
        )
        self.backup_id = self.client.post(reverse("ops-backups")).data["backup"]["id"]

    def _download(self, **params):
        return self.client.get(
            reverse("ops-backup-download", args=[self.backup_id]), params
        )

    def test_json_is_still_the_default_and_the_restorable_one(self):
        response = self._download()
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/json", response["Content-Type"])
        self.assertIn(".json", response["Content-Disposition"])

    def test_workbook_reads_in_arabic_with_names_instead_of_ids(self):
        from openpyxl import load_workbook

        response = self._download(format="xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames[0], "الملخص")
        self.assertIn("المنتجات", book.sheetnames)
        products = book["المنتجات"]
        self.assertEqual([c.value for c in products[1]][:3], ["الكود", "الاسم", "التصنيف"])
        row = [c.value for c in products[2]]
        self.assertEqual(row[0], "SKU1")
        self.assertEqual(row[2], "مواد غذائية")  # the category's name, not its id
        self.assertEqual(row[6], 6.5)            # money is a number a sheet can sum
        self.assertEqual(row[5], "0629000000001")  # a barcode keeps its leading zero
        self.assertEqual(row[-1], "نعم")          # true reads as a word
        invoices = book["الفواتير"]
        self.assertEqual([c.value for c in invoices[2]][1], "بقالة الحي")
        self.assertTrue(products.sheet_view.rightToLeft)

    def test_english_workbook_uses_english_headers(self):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(self._download(format="xlsx", lang="en").content))
        self.assertEqual(book.sheetnames[0], "Summary")
        self.assertEqual([c.value for c in book["Products"][1]][:2], ["SKU", "Name"])
        self.assertFalse(book["Products"].sheet_view.rightToLeft)

    def test_csv_zip_has_one_file_per_table_with_a_bom_for_excel(self):
        response = self._download(format="csv")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        self.assertIn(".zip", response["Content-Disposition"])
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        self.assertIn("المنتجات.csv", archive.namelist())
        text = archive.read("المنتجات.csv").decode("utf-8")
        self.assertTrue(text.startswith("﻿"))
        self.assertIn("أرز بسمتي", text)

    def test_an_unknown_format_is_refused(self):
        response = self._download(format="pdf")
        self.assertEqual(response.status_code, 400)
        self.assertIn("format", response.data)

    def test_another_company_cannot_read_the_file_in_any_format(self):
        other = Company.objects.create(name="Beta")
        role = Role.objects.get(name="Business Owner")
        User.objects.create_user(
            email="owner@beta.test", password="passw0rd123", company=other, role=role,
        )
        client = self.client_class()
        client.post(
            reverse("auth-login"),
            {"email": "owner@beta.test", "password": "passw0rd123", "device_id": "B"},
        )
        for fmt in ("json", "xlsx", "csv"):
            response = client.get(
                reverse("ops-backup-download", args=[self.backup_id]), {"format": fmt}
            )
            self.assertEqual(response.status_code, 404, fmt)
