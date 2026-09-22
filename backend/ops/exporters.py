"""Readable renderings of a backup snapshot.

The JSON dump is what a restore replays, and it is unreadable to the person
who owns the data: ids instead of names, raw timestamps, one long line. A
merchant who opens their backup wants to *see* their products, invoices and
payments — to check a figure, to hand a file to an accountant, or simply to
be sure the backup is really their shop.

So the same snapshot renders three ways:

- ``json``  — unchanged, the only format a restore accepts.
- ``xlsx``  — one sheet per table with translated headers, ids resolved to
  names from the dump's own master tables, and a summary sheet first.
- ``csv``   — a zip of the same sheets, UTF-8 with a BOM so Excel opens
  Arabic correctly, for anyone without Excel.

Everything is derived from the dump itself, never from the live database: a
backup taken in March must read as it was in March, not as things are now.
"""
import csv
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal

# Column spec: (key in the row, header key, optional lookup table name).
# A lookup turns "product": 19 into the product's name as it was at backup
# time; an id with no match falls back to "#19" rather than vanishing.
TABLES = [
    ("master", "products", "products", [
        ("sku", "sku", None), ("name", "name", None),
        ("category", "category", "categories"), ("brand", "brand", "brands"),
        ("unit", "unit", "units"), ("barcode", "barcode", None),
        ("cost_price", "cost", None), ("sale_price", "price", None),
        ("reorder_level", "reorderLevel", None), ("is_active", "active", None),
    ]),
    ("master", "customers", "customers", [
        ("name", "name", None), ("phone", "phone", None),
        ("email", "email", None), ("address", "address", None),
        ("is_active", "active", None),
    ]),
    ("master", "suppliers", "suppliers", [
        ("name", "name", None), ("phone", "phone", None),
        ("email", "email", None), ("address", "address", None),
        ("is_active", "active", None),
    ]),
    ("master", "categories", "categories", [
        ("name", "name", None), ("parent", "parent", "categories"),
        ("is_active", "active", None),
    ]),
    ("master", "brands", "brands", [("name", "name", None), ("is_active", "active", None)]),
    ("master", "units", "units", [
        ("name", "name", None), ("symbol", "symbol", None), ("is_active", "active", None),
    ]),
    ("master", "warehouses", "warehouses", [
        ("name", "name", None), ("code", "code", None), ("is_active", "active", None),
    ]),
    ("transactional", "invoices", "invoices", [
        ("number", "number", None), ("customer", "customer", "customers"),
        ("warehouse", "warehouse", "warehouses"), ("subtotal", "subtotal", None),
        ("tax_amount", "tax", None), ("total", "total", None),
        ("issued_at", "issuedAt", None),
    ]),
    ("transactional", "invoice_lines", "invoiceLines", [
        ("invoice", "invoice", "invoices"), ("product", "product", "products"),
        ("quantity", "quantity", None), ("unit_price", "price", None),
        ("line_subtotal", "subtotal", None), ("line_tax", "tax", None),
        ("line_total", "total", None),
    ]),
    ("transactional", "payments", "payments", [
        ("invoice", "invoice", "invoices"), ("method", "method", None),
        ("amount", "amount", None), ("recorded_at", "recordedAt", None),
    ]),
    ("transactional", "stock_movements", "stockMovements", [
        ("product", "product", "products"), ("warehouse", "warehouse", "warehouses"),
        ("movement_type", "movementType", None), ("quantity", "quantity", None),
        ("created_at", "createdAt", None),
    ]),
    ("transactional", "bills", "bills", [
        ("supplier", "supplier", "suppliers"), ("total", "total", None),
        ("is_void", "void", None), ("created_at", "createdAt", None),
    ]),
]

# Sheet titles and column headers, in the two languages the product speaks.
WORDS = {
    "ar": {
        "summary": "الملخص", "table": "الجدول", "rows": "عدد الصفوف",
        "company": "الشركة", "currency": "العملة", "takenAt": "وقت النسخة",
        "kind": "النوع", "note": "ملف للقراءة فقط — الاستعادة تحتاج ملف JSON.",
        "products": "المنتجات", "customers": "العملاء", "suppliers": "الموردون",
        "categories": "التصنيفات", "brands": "العلامات", "units": "الوحدات",
        "warehouses": "المستودعات", "invoices": "الفواتير", "invoiceLines": "بنود الفواتير",
        "payments": "المدفوعات", "stockMovements": "حركات المخزون", "bills": "فواتير الموردين",
        "sku": "الكود", "name": "الاسم", "category": "التصنيف", "brand": "العلامة",
        "unit": "الوحدة", "barcode": "الباركود", "cost": "التكلفة", "price": "السعر",
        "reorderLevel": "حد إعادة الطلب", "active": "نشط", "phone": "الهاتف",
        "email": "البريد", "address": "العنوان", "parent": "التصنيف الأب",
        "symbol": "الرمز", "code": "الرمز", "number": "الرقم", "customer": "العميل",
        "warehouse": "المستودع", "subtotal": "الإجمالي قبل الضريبة", "tax": "الضريبة",
        "total": "الإجمالي", "issuedAt": "تاريخ الإصدار", "invoice": "الفاتورة",
        "product": "المنتج", "quantity": "الكمية", "method": "طريقة الدفع",
        "amount": "المبلغ", "recordedAt": "وقت التسجيل", "movementType": "نوع الحركة",
        "createdAt": "وقت الإنشاء", "supplier": "المورد", "void": "ملغاة",
        "yes": "نعم", "no": "لا",
    },
    "en": {
        "summary": "Summary", "table": "Table", "rows": "Rows",
        "company": "Company", "currency": "Currency", "takenAt": "Taken at",
        "kind": "Kind", "note": "A read-only copy — a restore needs the JSON file.",
        "products": "Products", "customers": "Customers", "suppliers": "Suppliers",
        "categories": "Categories", "brands": "Brands", "units": "Units",
        "warehouses": "Warehouses", "invoices": "Invoices", "invoiceLines": "Invoice lines",
        "payments": "Payments", "stockMovements": "Stock movements", "bills": "Supplier bills",
        "sku": "SKU", "name": "Name", "category": "Category", "brand": "Brand",
        "unit": "Unit", "barcode": "Barcode", "cost": "Cost", "price": "Price",
        "reorderLevel": "Reorder level", "active": "Active", "phone": "Phone",
        "email": "Email", "address": "Address", "parent": "Parent category",
        "symbol": "Symbol", "code": "Code", "number": "Number", "customer": "Customer",
        "warehouse": "Warehouse", "subtotal": "Subtotal", "tax": "Tax",
        "total": "Total", "issuedAt": "Issued at", "invoice": "Invoice",
        "product": "Product", "quantity": "Quantity", "method": "Payment method",
        "amount": "Amount", "recordedAt": "Recorded at", "movementType": "Movement type",
        "createdAt": "Created at", "supplier": "Supplier", "void": "Void",
        "yes": "Yes", "no": "No",
    },
}

FORMATS = ("json", "xlsx", "csv")

# Columns that are money or quantities: their values become real numbers so a
# spreadsheet can sum them. Everything else stays text — a barcode, a phone or
# an SKU with a leading zero must not be "helpfully" turned into a number.
NUMERIC = frozenset({
    "cost", "price", "subtotal", "tax", "total", "amount", "quantity", "reorderLevel",
})


def _words(language):
    return WORDS.get(language, WORDS["ar"])


def _lookups(dump):
    """id → label per master table, from the snapshot itself."""
    maps = {}
    for section in ("master", "transactional"):
        for table, rows in (dump.get(section) or {}).items():
            label = "number" if table == "invoices" else "name"
            maps[table] = {
                row.get("id"): row.get(label) for row in rows if row.get("id") is not None
            }
    return maps


def _cell(value, words, numeric=False):
    """One value as a person reads it: dates without the T and the Z, yes/no
    instead of true/false, an empty cell instead of the word None.

    A dump read back from storage carries strings (it went through JSON); one
    built in memory carries real datetimes and Decimals. Both arrive here, and
    a tz-aware datetime has to lose its tzinfo — Excel cannot store one.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return words["yes" if value else "no"]
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Decimal):
        return float(value)
    if numeric and isinstance(value, str) and value:
        try:
            return float(value)
        except ValueError:
            return value
    if isinstance(value, str) and len(value) >= 19 and value[10:11] == "T":
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            return parsed.strftime("%Y-%m-%d %H:%M")
        except ValueError:
            return value
    return value


def sheets(dump, language="ar"):
    """[(title, [header row], [data rows])] — the shape both writers render."""
    words = _words(language)
    maps = _lookups(dump)
    out = []
    for section, table, title_key, columns in TABLES:
        rows = (dump.get(section) or {}).get(table) or []
        header = [words.get(header_key, header_key) for _key, header_key, _l in columns]
        body = []
        for row in rows:
            line = []
            for key, header_key, lookup in columns:
                value = row.get(key)
                if lookup and value is not None:
                    value = maps.get(lookup, {}).get(value) or f"#{value}"
                    if not isinstance(value, str):
                        value = str(value)
                line.append(_cell(value, words, numeric=header_key in NUMERIC))
            body.append(line)
        out.append((words.get(title_key, title_key), header, body))
    return out


def _summary(dump, record, language):
    words = _words(language)
    company = dump.get("company") or {}
    header = [words["table"], words["rows"]]
    rows = [[title, len(body)] for title, _h, body in sheets(dump, language)]
    facts = [
        [words["company"], company.get("name") or ""],
        [words["currency"], company.get("currency") or ""],
        [words["takenAt"], record.created_at.strftime("%Y-%m-%d %H:%M") if record else ""],
        [words["kind"], getattr(record, "kind", "") or ""],
        [words["note"], ""],
        ["", ""],
    ]
    return words["summary"], header, facts + rows


def to_xlsx(dump, record=None, language="ar"):
    """The whole snapshot as one workbook: summary first, then a sheet per
    table. Returns the file's bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    rtl = language != "en"
    for title, header, body in [_summary(dump, record, language)] + sheets(dump, language):
        sheet = book.create_sheet(title[:31])
        sheet.sheet_view.rightToLeft = rtl
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="right" if rtl else "left")
        for line in body:
            sheet.append(line)
        sheet.freeze_panes = "A2"
        for index, name in enumerate(header, start=1):
            widest = max(
                [len(str(name))] + [len(str(line[index - 1])) for line in body[:200]] or [0]
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 10), 46)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_csv_zip(dump, record=None, language="ar"):
    """One CSV per table inside a zip. The BOM matters: without it Excel
    reads Arabic as mojibake, which is exactly the audience for this format."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for title, header, body in [_summary(dump, record, language)] + sheets(dump, language):
            text = io.StringIO()
            writer = csv.writer(text)
            writer.writerow(header)
            writer.writerows(body)
            archive.writestr(f"{title}.csv", "﻿" + text.getvalue())
    return buffer.getvalue()


RENDERERS = {
    "xlsx": (
        to_xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    "csv": (to_csv_zip, "application/zip", "zip"),
}
